"""Genera el nuevo modelo DIM_HOMOLOGADOS (v2): Excel + scripts de carga SQL Server 2012.

Ejecutar desde la raíz del repo:  python dim_homologados_v2/scripts/generar_dim_homologados.py

Entradas:
  items/backup/dim_homologados_minsa_post_script3_20261006.xlsx  base vigente (con CodigoSIGA)
  items/items.csv                    maestro de productos Unilene (Coditem, CodSut, ...)
  items/items_minsa.csv              catálogo MINSA (CodigoMed, NombreMed)
  items/items_vendidos_minsa.xlsx    SKU vendidos a MINSA
  salida/homologacion_siga_final.xlsx
  salida/dim_homologados_minsa_siga.xlsx (hojas Revision_manual y Revision_vendidos)
Salidas:
  dim_homologados_v2/dim_homologados_minsa_v2.xlsx
  dim_homologados_v2/sql/02_..., 03_..., 04_... (datos)
  items/dim_homologados_minsa.xlsx (plano, actualizado y depurado)
"""
import re
import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path('.')
OUT = ROOT / 'dim_homologados_v2'
HOY = datetime.date.today()
BASE = ROOT / 'items/backup/dim_homologados_minsa_post_script3_20261006.xlsx'
LEGACY = ['ID_SKU_EXTERNO', 'DESC_PRODUCTO_EXTERNO', 'FUENTE', 'ID_SKU_UNICO_UNILENE', 'ID_SKU_UNILENE', 'AGRUPADOR',
          'SUBAGRUPADOR', 'AGRUPADOR_NEWCOMERCIAL', 'SUB_AGRUPADOR_NEWCOMERCIAL', 'LINEA_NEWCOMERCIAL',
          'DPRO_DESC_SUB_FAMILIA']
RX_DOBLE = re.compile(r'C/\s?2\s?A\b|DOBLE AGUJA|\b2C/A|C/2 AGUJAS')


def nz(x):
    if x is None or (isinstance(x, float) and np.isnan(x)) or str(x).strip() == '':
        return None
    return str(x).strip()


# ---------------------------------------------------------------- maestro Unilene
it = pd.read_csv(ROOT / 'items/items.csv', dtype=str).apply(lambda c: c.str.strip())
sut = (it.Agrupador == '01. SUTURAS') & (it.CodSut.str.len() == 21)
it['ES_SUTURA'] = sut
it['DOBLE'] = sut & (it.CodSut.str[17] == 'B')
it['UNICO'] = np.where(sut, it.CodSut.str[2:15] + np.where(it.DOBLE, 'B', ''), it.Coditem)
it['SUBAGR'] = np.where(sut, 'SUTURAS', it.SubFamilia)
I = it.set_index('Coditem')

# ---------------------------------------------------------------- catálogo MINSA, vendidos, SIGA, base
def leer_catalogo_minsa(path):
    """items_minsa.csv mezcla lineas UTF-8 y Windows-1252: se decodifica linea por linea."""
    import csv
    import io
    lineas = []
    for raw in path.read_bytes().splitlines():
        try:
            lineas.append(raw.decode('utf-8'))
        except UnicodeDecodeError:
            lineas.append(raw.decode('cp1252', errors='replace'))
    df = pd.DataFrame(list(csv.reader(io.StringIO('\n'.join(lineas))))[1:], columns=['CodigoMed', 'NombreMed'])
    df['CodigoMed'] = df.CodigoMed.str.strip()
    df['NombreMed'] = df.NombreMed.map(lambda s: re.sub(r'[\x00-\x1f\x7f]', ' ', s).strip())
    return df.drop_duplicates('CodigoMed').set_index('CodigoMed').NombreMed


mi = leer_catalogo_minsa(ROOT / 'items/items_minsa.csv')
vend = set(pd.read_excel(ROOT / 'items/items_vendidos_minsa.xlsx', header=None, dtype=str)[0].str.strip())
sg = pd.read_excel(ROOT / 'salida/homologacion_siga_final.xlsx', dtype=str).apply(lambda c: c.str.strip()).set_index('CodigoMed')
d = pd.read_excel(BASE, dtype=str).apply(lambda c: c.str.strip())
N_BASE = len(d)
dep = []


def depurar(rows, motivo):
    for _, r in rows.iterrows():
        dep.append({**{c: r.get(c) for c in LEGACY}, 'MOTIVO_DEPURACION': motivo})


# 1) SKU inexistente en el maestro
d['EN_ITEMS'] = d.ID_SKU_UNILENE.isin(I.index)
d['VENDIDO'] = d.ID_SKU_UNILENE.isin(vend)
d['VALIDACION'] = 'OK'
d['MOTIVO_REVISION'] = None
m = ~d.EN_ITEMS & ~d.VENDIDO
depurar(d[m], 'SKU no existe en el maestro items.csv (carga antigua)')
d = d[~m].copy()
m = ~d.EN_ITEMS
d.loc[m, 'VALIDACION'] = 'REVISION'
d.loc[m, 'MOTIVO_REVISION'] = 'SKU vendido pero no existe en items.csv'

# 2) datos maestros desde items.csv (ID único con sufijo B de doble aguja)
for c, src in [('ID_SKU_UNICO_UNILENE', 'UNICO'), ('AGRUPADOR', 'Agrupador'), ('SUBAGRUPADOR', 'SUBAGR'),
               ('DPRO_DESC_SUB_FAMILIA', 'SubFamilia')]:
    d.loc[d.EN_ITEMS, c] = d.loc[d.EN_ITEMS, 'ID_SKU_UNILENE'].map(I[src])

# 3) descripción oficial del catálogo MINSA; si no está, la más completa
alt = d.groupby('ID_SKU_EXTERNO').DESC_PRODUCTO_EXTERNO.agg(
    lambda x: max(x.dropna(), key=lambda s: (s.startswith('SUTURA') or 'UNIDAD' in s, len(s))))
d['DESC_PRODUCTO_EXTERNO'] = d.ID_SKU_EXTERNO.map(mi).fillna(d.ID_SKU_EXTERNO.map(alt))

# 4) consistencia doble aguja: descripción MINSA vs CodSut
d['DOBLE_MINSA'] = d.DESC_PRODUCTO_EXTERNO.str.contains(RX_DOBLE)
d['DOBLE_SKU'] = d.ID_SKU_UNILENE.map(I.DOBLE)
es_sut = d.ID_SKU_UNILENE.map(I.ES_SUTURA).astype('boolean').fillna(False).astype(bool)
inc = d.EN_ITEMS & es_sut & (d.DOBLE_MINSA != d.DOBLE_SKU)
m = inc & ~d.VENDIDO
depurar(d[m], 'Inconsistencia doble aguja: descripción MINSA y CodSut no coinciden (SKU no vendido)')
d = d[~m].copy()
m = inc.loc[d.index]
d.loc[m, 'VALIDACION'] = 'REVISION'
d.loc[m, 'MOTIVO_REVISION'] = 'Inconsistencia doble aguja con SKU vendido: validar'

# 4b) corrección: para esos códigos MINSA se agrega la familia correcta (mismo ID único con/sin "B", marcas SD/CQ)
it['MARCA'] = np.where(sut, it.CodSut.str[:2], '')
fam_ok = it[~it.ES_SUTURA | it.MARCA.isin(['SD', 'CQ'])].groupby('UNICO').Coditem.apply(list)
ya = set(zip(d.ID_SKU_EXTERNO, d.ID_SKU_UNILENE))
nuevas = []
for r in [x for x in dep if x['MOTIVO_DEPURACION'].startswith('Inconsistencia doble aguja')]:
    ext = r['ID_SKU_EXTERNO']
    desc = d.loc[d.ID_SKU_EXTERNO == ext, 'DESC_PRODUCTO_EXTERNO']
    desc = desc.iloc[0] if len(desc) else mi.get(ext, r['DESC_PRODUCTO_EXTERNO'])
    doble = bool(RX_DOBLE.search(str(desc)))
    u = I.UNICO.get(r['ID_SKU_UNILENE'])
    if not isinstance(u, str):
        continue
    target = u.rstrip('B') + ('B' if doble else '')
    for cod in fam_ok.get(target, []):
        if (ext, cod) in ya:
            continue
        ya.add((ext, cod))
        nuevas.append({'ID_SKU_EXTERNO': ext, 'DESC_PRODUCTO_EXTERNO': desc, 'FUENTE': 'MINSA', 'ID_SKU_UNICO_UNILENE': target,
                       'ID_SKU_UNILENE': cod, 'AGRUPADOR': I.Agrupador[cod], 'SUBAGRUPADOR': I.SUBAGR[cod],
                       'DPRO_DESC_SUB_FAMILIA': I.SubFamilia[cod], 'EN_ITEMS': True, 'VENDIDO': cod in vend,
                       'VALIDACION': 'OK', 'MOTIVO_REVISION': None, 'ORIGEN_FILA': 'Corrección doble aguja (CodSut)'})
CORR = pd.DataFrame(nuevas)
if len(CORR):
    d = pd.concat([d, CORR], ignore_index=True)

# 5) medida incoherente conocida (30005: MINSA 40 mm, SKU 30 mm)
m = (d.ID_SKU_EXTERNO == '30005') & (d.ID_SKU_UNICO_UNILENE == 'CC030MR300075')
depurar(d[m], 'Medida de aguja distinta a la descripción MINSA (30 mm vs 40 mm)')
d = d[~m].copy()

# 6) códigos MINSA duplicados (mismo producto y mismos SKU) -> código canónico
def norm(s):
    return re.sub(r'\s', '', re.sub(r'\*|\bUNIDAD\b', '', str(s).upper()))


P = d.groupby('ID_SKU_EXTERNO').agg(DESC=('DESC_PRODUCTO_EXTERNO', 'first'),
                                     SKUS=('ID_SKU_UNILENE', lambda x: frozenset(x))).reset_index()
P['SIGA_ESTADO'] = P.ID_SKU_EXTERNO.map(sg.SIGA_estado)
P['SIGA_COD'] = P.ID_SKU_EXTERNO.map(sg.SIGA_codigo_1)
P['K_DESC'] = P.DESC.map(norm)
parent = {e: e for e in P.ID_SKU_EXTERNO}


def find(x):
    while parent[x] != x:
        parent[x] = parent[parent[x]]
        x = parent[x]
    return x


for key in ['SIGA_COD', 'K_DESC']:
    sub = P[P[key].notna() & ((P.SIGA_ESTADO == 'exacto') if key == 'SIGA_COD' else True)]
    for _, g in sub.groupby([key, 'SKUS']):
        e = list(g.ID_SKU_EXTERNO)
        for x in e[1:]:
            parent[find(x)] = find(e[0])
P['GRUPO'] = P.ID_SKU_EXTERNO.map(find)
P['R'] = P.apply(lambda r: (0 if r.ID_SKU_EXTERNO in mi.index else 1, 1 if '*' in str(r.DESC) else 0,
                            0 if r.SIGA_ESTADO == 'exacto' else 1,
                            int(r.ID_SKU_EXTERNO) if r.ID_SKU_EXTERNO.isdigit() else 99999), axis=1)
can = P.sort_values('R').groupby('GRUPO').ID_SKU_EXTERNO.first()
canon = P.set_index('ID_SKU_EXTERNO').GRUPO.map(can)
m = d.ID_SKU_EXTERNO.map(canon) != d.ID_SKU_EXTERNO
for _, r in d[m].iterrows():
    dep.append({**{c: r.get(c) for c in LEGACY},
                'MOTIVO_DEPURACION': f'Código MINSA duplicado: mismo producto que {canon[r.ID_SKU_EXTERNO]} (canónico)'})
d = d[~m].copy()

# 7) atributos finales
d['NIVEL'] = 'EXACTO'
d['ORIGEN'] = d.get('ORIGEN_FILA', pd.Series(index=d.index, dtype=object)).fillna('DIM_HOMOLOGADOS + homologacion_siga_final')
d = d.sort_values(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE']).reset_index(drop=True)
assert not d.duplicated(['FUENTE', 'ID_SKU_EXTERNO', 'ID_SKU_UNILENE']).any()
DEP = pd.DataFrame(dep)

# ---------------------------------------------------------------- HOMOLOGACION_REVISION
an = ROOT / 'salida/dim_homologados_minsa_siga.xlsx'
rm = pd.read_excel(an, sheet_name='Revision_manual', dtype=str)
rv = pd.read_excel(an, sheet_name='Revision_vendidos', dtype=str)
RC = ['ID_SKU_EXTERNO', 'ID_SKU_UNILENE', 'TIER_COL_F', 'COL', 'PRIORIDAD', 'SCORE', 'DIFERENCIAS', 'ORIGEN']
REV = pd.concat([
    rm.assign(ORIGEN='siga_final F=probable/revisar', COL=rm.COL_ORIGEN)[RC],
    rv.assign(ORIGEN='rescate SKU vendidos', PRIORIDAD=None)[RC]]).apply(
    lambda c: c.str.strip() if c.dtype == object else c)
DBL = DEP[DEP.MOTIVO_DEPURACION.str.startswith('Inconsistencia doble aguja')]
REV = pd.concat([REV, pd.DataFrame({'ID_SKU_EXTERNO': DBL.ID_SKU_EXTERNO, 'ID_SKU_UNILENE': DBL.ID_SKU_UNILENE,
                                    'TIER_COL_F': 'base_dim', 'COL': None, 'PRIORIDAD': None, 'SCORE': None,
                                    'DIFERENCIAS': 'doble aguja: MINSA vs CodSut',
                                    'ORIGEN': 'depurado doble aguja (validar)'})], ignore_index=True)
REV['ID_SKU_EXTERNO'] = REV.ID_SKU_EXTERNO.map(lambda e: canon.get(e, e))
REV = REV[REV.ID_SKU_UNILENE.isin(I.index)]
REV = REV[~(REV.ID_SKU_EXTERNO + '|' + REV.ID_SKU_UNILENE).isin(set(d.ID_SKU_EXTERNO + '|' + d.ID_SKU_UNILENE))]
REV = REV.sort_values(['ID_SKU_EXTERNO', 'PRIORIDAD']).drop_duplicates(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE'])
REV['ID_SKU_UNICO_UNILENE'] = REV.ID_SKU_UNILENE.map(I.UNICO)
REV['ITEM_UNILENE'] = REV.ID_SKU_UNILENE.map(I.Item)
REV['VENDIDO_MINSA'] = np.where(REV.ID_SKU_UNILENE.isin(vend), 'SI', 'NO')
REV['FUENTE'] = 'MINSA'
REV['DECISION'] = None
REVCOLS = ['FUENTE', 'ID_SKU_EXTERNO', 'ID_SKU_UNILENE', 'ID_SKU_UNICO_UNILENE', 'ITEM_UNILENE', 'TIER_COL_F', 'COL',
           'PRIORIDAD', 'SCORE', 'DIFERENCIAS', 'VENDIDO_MINSA', 'ORIGEN', 'DECISION']
REV = REV[REVCOLS].reset_index(drop=True)

# ---------------------------------------------------------------- DIM_PRODUCTO_EXTERNO
codes = sorted(set(d.ID_SKU_EXTERNO) | set(P.ID_SKU_EXTERNO) | set(REV.ID_SKU_EXTERNO))
descs = {**sg.NombreMed.to_dict(), **P.set_index('ID_SKU_EXTERNO').DESC.to_dict(), **mi.to_dict()}
PR = pd.DataFrame({'ID_SKU_EXTERNO': codes})
PR['FUENTE'] = 'MINSA'
PR['DESC_PRODUCTO_EXTERNO'] = PR.ID_SKU_EXTERNO.map(descs)
PR['ID_SKU_EXTERNO_CANONICO'] = PR.ID_SKU_EXTERNO.map(lambda e: canon.get(e, e))
PR['ES_DUPLICADO'] = (PR.ID_SKU_EXTERNO != PR.ID_SKU_EXTERNO_CANONICO).astype(int)
PR['EN_CATALOGO_MINSA'] = PR.ID_SKU_EXTERNO.isin(mi.index).astype(int)
PR['SIGA_ESTADO'] = PR.ID_SKU_EXTERNO.map(sg.SIGA_estado).fillna('sin_busqueda')
PR['CodigoSIGA'] = np.where(PR.SIGA_ESTADO == 'exacto', PR.ID_SKU_EXTERNO.map(sg.SIGA_codigo_1), None)
PR['CodigoSIGA_SUGERIDO'] = np.where(PR.SIGA_ESTADO != 'exacto', PR.ID_SKU_EXTERNO.map(sg.SIGA_codigo_1), None)
PR['DESC_SIGA'] = PR.ID_SKU_EXTERNO.map(sg.SIGA_descripcion_1)
# el canónico hereda el CodigoSIGA exacto de su duplicado si no lo tiene
for g, grp in PR.groupby('ID_SKU_EXTERNO_CANONICO'):
    if len(grp) > 1:
        ci = grp.index[grp.ID_SKU_EXTERNO == g]
        cs = grp.CodigoSIGA.dropna()
        if len(ci) and pd.isna(PR.loc[ci[0], 'CodigoSIGA']) and len(cs):
            PR.loc[ci[0], 'CodigoSIGA'] = cs.iloc[0]
PR['VALIDACION_SIGA'] = np.where(PR.CodigoSIGA.notna(), 'OK', np.where(PR.SIGA_ESTADO == 'sin_busqueda', 'SIN_DATO', 'REVISION'))
PR['TIENE_HOMOLOGACION'] = PR.ID_SKU_EXTERNO.isin(set(d.ID_SKU_EXTERNO)).astype(int)
PRCOLS = ['FUENTE', 'ID_SKU_EXTERNO', 'DESC_PRODUCTO_EXTERNO', 'ID_SKU_EXTERNO_CANONICO', 'ES_DUPLICADO',
          'EN_CATALOGO_MINSA', 'CodigoSIGA', 'VALIDACION_SIGA', 'SIGA_ESTADO', 'CodigoSIGA_SUGERIDO', 'DESC_SIGA',
          'TIENE_HOMOLOGACION']
PR = PR[PRCOLS]
d['CodigoSIGA'] = d.ID_SKU_EXTERNO.map(PR.set_index('ID_SKU_EXTERNO').CodigoSIGA)
DCOLS = LEGACY + ['CodigoSIGA', 'VALIDACION', 'MOTIVO_REVISION', 'NIVEL', 'ORIGEN']

# ---------------------------------------------------------------- Excel plano (base de trabajo)
wb = Workbook()
ws = wb.active
ws.title = 'sheet1'
pc = LEGACY + ['CodigoSIGA', 'VALIDACION']
ws.append(pc)
for r in d[pc].itertuples(index=False):
    vals = [nz(x) for x in r]
    if vals[0] and vals[0].isdigit() and not vals[0].startswith('0'):
        vals[0] = int(vals[0])
    ws.append(vals)
wb.save(ROOT / 'items/dim_homologados_minsa.xlsx')

# ---------------------------------------------------------------- Excel v2
motivos = DEP.MOTIVO_DEPURACION.str.replace(r'que \d+ \(canónico\)', 'que otro código (canónico)', regex=True).value_counts()
RES = pd.DataFrame(
    [('DIM_HOMOLOGADOS v2 (MINSA)', None, None),
     ('Filas base anterior', N_BASE, 'items/backup/dim_homologados_minsa_post_script3_20261006.xlsx'),
     ('Filas depuradas (ruido)', -len(DEP), 'Hoja DEPURADOS'),
     ('Filas agregadas por corrección doble aguja', len(CORR), 'Familia correcta según CodSut (SD/CQ)'),
     ('Filas finales', len(d), None),
     ('  VALIDACION = OK', int((d.VALIDACION == 'OK').sum()), None),
     ('  VALIDACION = REVISION', int((d.VALIDACION == 'REVISION').sum()), 'Vigentes (SKU vendidos) pero con inconsistencia'),
     ('Códigos MINSA con homologación', d.ID_SKU_EXTERNO.nunique(), None),
     ('  con CodigoSIGA', int(d.drop_duplicates('ID_SKU_EXTERNO').CodigoSIGA.notna().sum()), None),
     ('SKU vendidos a MINSA cubiertos', len(set(d.ID_SKU_UNILENE) & vend), f'de {len(vend)}'),
     (None, None, None), ('DEPURACIÓN', None, None)]
    + [(f'  {k}', v, None) for k, v in motivos.items()]
    + [(None, None, None), ('DIM_PRODUCTO_EXTERNO', len(PR), None),
       ('  códigos duplicados mapeados a canónico', int(PR.ES_DUPLICADO.sum()), 'Columna ID_SKU_EXTERNO_CANONICO'),
       ('  fuera del catálogo MINSA actual', int((PR.EN_CATALOGO_MINSA == 0).sum()), None),
       ('  VALIDACION_SIGA = OK / REVISION / SIN_DATO',
        ' / '.join(str(int((PR.VALIDACION_SIGA == k).sum())) for k in ('OK', 'REVISION', 'SIN_DATO')), None),
       (None, None, None),
       ('HOMOLOGACION_REVISION (candidatos pendientes)', len(REV),
        f'{REV.ID_SKU_EXTERNO.nunique()} códigos MINSA; {int((REV.VENDIDO_MINSA == "SI").sum())} filas con SKU vendido')],
    columns=['Concepto', 'Valor', 'Nota'])
LEEME = pd.DataFrame({'Regla / definición': [
    'Clave de DIM_HOMOLOGADOS: (FUENTE, ID_SKU_EXTERNO, ID_SKU_UNILENE). Sin duplicados.',
    'ID_SKU_UNICO_UNILENE (suturas) = CodSut posiciones 3-15 + "B" si es doble aguja (posición 18 = B). Otros agrupadores = Coditem.',
    'Familia homologable: todos los Coditem con el mismo ID_SKU_UNICO y marca SD o CQ.',
    'AGRUPADOR, SUBAGRUPADOR y DPRO_DESC_SUB_FAMILIA se toman siempre del maestro items.csv.',
    'DESC_PRODUCTO_EXTERNO = descripción oficial del catálogo MINSA (items_minsa.csv).',
    'Códigos MINSA duplicados (mismo producto / mismo código SIGA y mismos SKU): se conserva un canónico '
    '(en catálogo > sin * > SIGA exacto > código menor). Los demás quedan en DIM_PRODUCTO_EXTERNO apuntando al canónico.',
    'CodigoSIGA vive en DIM_PRODUCTO_EXTERNO (atributo del producto MINSA). Se asigna solo si la búsqueda SIGA fue exacta.',
    'VALIDACION: OK = homologación aprobada; REVISION = vigente pero requiere validar; RECHAZADO = no usar.',
    'HOMOLOGACION_REVISION: candidatos aproximados. Completar DECISION (SI/NO) y ejecutar dbo.usp_HOMOLOGACION_APROBAR_REVISION.',
    'Para hechos (ventas/compras MINSA) usar dbo.VW_HOMOLOGACION_RESUELTA: resuelve cualquier código, incluso duplicados.']})
sheets = [('RESUMEN', RES), ('LEEME', LEEME), ('DIM_HOMOLOGADOS', d[DCOLS]), ('DIM_PRODUCTO_EXTERNO', PR),
          ('HOMOLOGACION_REVISION', REV), ('DEPURADOS', DEP)]
hf = PatternFill('solid', fgColor='1F4E78')
sec = PatternFill('solid', fgColor='D9E1F2')
fills = {'OK': 'C6EFCE', 'REVISION': 'FFF2CC', 'SIN_DATO': 'F8CBAD'}
with pd.ExcelWriter(OUT / 'dim_homologados_minsa_v2.xlsx', engine='openpyxl') as w:
    for n, df in sheets:
        st = 2 if n == 'RESUMEN' else 0
        df.to_excel(w, sheet_name=n, index=False, startrow=st)
        ws = w.sheets[n]
        h = st + 1
        for c in ws[h]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = hf
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for i, cn in enumerate(df.columns, 1):
            L = max([len(str(cn))] + [len(str(x)) for x in df[cn].head(500).fillna('')])
            ws.column_dimensions[get_column_letter(i)].width = min(max(11, L + 2), 140 if n == 'LEEME' else 60)
        if n == 'RESUMEN':
            ws['A1'] = 'DIM_HOMOLOGADOS v2 – MINSA'
            ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
            ws['A2'] = f'Generado {HOY:%d/%m/%Y}'
            ws['A2'].font = Font(italic=True, color='595959')
            for r in range(h + 1, ws.max_row + 1):
                a = ws.cell(r, 1).value
                if a and not str(a).startswith(' ') and ws.cell(r, 2).value in (None, '') or (a and str(a).split(' (')[0].isupper()):
                    for j in (1, 2, 3):
                        ws.cell(r, j).fill = sec
                        ws.cell(r, j).font = Font(bold=True, color='1F4E78')
            ws.column_dimensions['A'].width = 70
            ws.column_dimensions['C'].width = 55
            continue
        if n == 'LEEME':
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
            continue
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        for colname in ('VALIDACION', 'VALIDACION_SIGA'):
            if colname in df.columns:
                j = list(df.columns).index(colname) + 1
                for r in range(2, ws.max_row + 1):
                    f = fills.get(ws.cell(r, j).value)
                    if f:
                        ws.cell(r, j).fill = PatternFill('solid', fgColor=f)


# ---------------------------------------------------------------- SQL de datos (SQL Server 2012)
def q(x):
    if isinstance(x, (int, np.integer)):
        return str(int(x))
    x = nz(x)
    return 'NULL' if x is None else "N'" + x.replace("'", "''") + "'"


def inserts(tabla, df, cols, lote=500):
    out = []
    for i in range(0, len(df), lote):
        chunk = df.iloc[i:i + lote]
        out.append(f"INSERT INTO {tabla} ({', '.join(cols)}) VALUES")
        out.append(',\n'.join('  (' + ', '.join(q(getattr(r, c)) for c in cols) + ')'
                              for r in chunk[cols].itertuples(index=False)) + ';')
    return out


def w_sql(name, lines):
    (OUT / 'sql' / name).write_text('\n'.join(lines) + '\n', encoding='utf-8-sig')


def head(t):
    return [f'/* {t}', f'   DIM_HOMOLOGADOS v2 – generado {HOY} – compatible con SQL Server 2012', '*/',
            'SET NOCOUNT ON;', 'SET XACT_ABORT ON;', 'GO', '']


w_sql('02_cargar_dim_producto_externo.sql', head('02 – Carga DIM_PRODUCTO_EXTERNO (MINSA)') + [
    'BEGIN TRANSACTION;', "DELETE FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = 'MINSA';",
    "DELETE FROM dbo.DIM_PRODUCTO_EXTERNO WHERE FUENTE = 'MINSA';", '']
      + inserts('dbo.DIM_PRODUCTO_EXTERNO', PR, PRCOLS)
      + ['', f'-- Esperado: {len(PR)} filas MINSA',
         "SELECT COUNT(*) AS filas FROM dbo.DIM_PRODUCTO_EXTERNO WHERE FUENTE = 'MINSA';", 'COMMIT TRANSACTION;', 'GO'])
HC = LEGACY + ['VALIDACION', 'MOTIVO_REVISION', 'NIVEL', 'ORIGEN']
w_sql('03_cargar_dim_homologados.sql', head('03 – Carga DIM_HOMOLOGADOS (MINSA) depurado') + [
    'BEGIN TRANSACTION;', "DELETE FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = 'MINSA';", '']
      + inserts('dbo.DIM_HOMOLOGADOS', d, HC)
      + ['', f"-- Esperado: {len(d)} filas MINSA ({int((d.VALIDACION == 'OK').sum())} OK / {int((d.VALIDACION == 'REVISION').sum())} REVISION)",
         "SELECT VALIDACION, COUNT(*) AS filas FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = 'MINSA' GROUP BY VALIDACION;",
         'COMMIT TRANSACTION;', 'GO'])
w_sql('04_cargar_homologacion_revision.sql', head('04 – Carga HOMOLOGACION_REVISION (candidatos pendientes)') + [
    'BEGIN TRANSACTION;', "DELETE FROM dbo.HOMOLOGACION_REVISION WHERE FUENTE = 'MINSA' AND DECISION IS NULL;", '']
      + inserts('dbo.HOMOLOGACION_REVISION', REV, REVCOLS[:-1])
      + ['', f'-- Esperado: {len(REV)} filas',
         "SELECT COUNT(*) AS filas FROM dbo.HOMOLOGACION_REVISION WHERE FUENTE = 'MINSA';", 'COMMIT TRANSACTION;', 'GO'])

STG = it[['Coditem', 'CodSut', 'Item', 'Agrupador', 'SubFamilia']].drop_duplicates('Coditem')
w_sql('08_cargar_stg_items.sql', head('08 – Carga STG_ITEMS (maestro Unilene desde items.csv). Reemplazar luego por carga desde el ERP') + [
    'BEGIN TRANSACTION;', 'TRUNCATE TABLE dbo.STG_ITEMS;', '']
      + inserts('dbo.STG_ITEMS', STG, list(STG.columns))
      + ['', f'-- Esperado: {len(STG)} filas', 'SELECT COUNT(*) AS filas FROM dbo.STG_ITEMS;', 'COMMIT TRANSACTION;', 'GO'])

print(RES.to_string())
print('max len:', {c: int(d[c].dropna().str.len().max()) for c in HC if d[c].notna().any()})
print('max len PR:', {c: int(PR[c].dropna().astype(str).str.len().max()) for c in PRCOLS if PR[c].notna().any()})
print('max len REV:', {c: int(REV[c].dropna().astype(str).str.len().max()) for c in REVCOLS if REV[c].notna().any()})
