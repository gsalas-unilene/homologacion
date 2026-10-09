"""Aplica la validación automática de suturas sobre DIM_HOMOLOGADOS (parte 6).

Ejecutar desde la raíz:  python scripts/aplicar_validacion_auto.py
1) Filas actuales con nº de agujas distinto (o medida 30005): se eliminan y se busca su homólogo correcto en items.csv.
2) Candidatos de siga_final (hoja 1B) de suturas: AUTO_SI se agregan, AUTO_NO se descartan, MANUAL quedan.
3) SKU vendidos sin homologación (hoja 1C) de suturas: se busca en el catálogo MINSA el código AUTO_SI.
4) Se revisa toda la tabla actual con el validador y se reportan alertas (no se borran).
Salidas: items/dim_homologados_minsa.xlsx, salida/validacion_auto_dim_homologados.sql,
         salida/validacion_auto_resultados.xlsx, salida/decisiones_auto.csv
"""
import csv
import datetime
import io
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import validador_auto as V  # noqa: E402

ROOT = Path('.')
HOY = datetime.date.today()
BASE = ROOT / 'items/backup/dim_homologados_minsa_post_script5_20261007.xlsx'
SRC = ROOT / 'items/dim_homologados_minsa.xlsx'
if not BASE.exists():
    BASE.write_bytes(SRC.read_bytes())

# ------------------------------------------------------------------ datos
it = pd.read_csv(ROOT / 'items/items.csv', dtype=str).apply(lambda c: c.str.strip())
sut = (it.Agrupador == '01. SUTURAS') & (it.CodSut.str.len() == 21)
it['UNICO'] = np.where(sut, it.CodSut.str[2:15], it.Coditem)
it['MARCA'] = np.where(sut, it.CodSut.str[:2], '')
I = it.set_index('Coditem')
vend = set(pd.read_excel(ROOT / 'items/items_vendidos_minsa.xlsx', header=None, dtype=str)[0].str.strip())


def leer_catalogo(path):
    lineas = []
    for raw in path.read_bytes().splitlines():
        try:
            lineas.append(raw.decode('utf-8'))
        except UnicodeDecodeError:
            lineas.append(raw.decode('cp1252', errors='replace'))
    df = pd.DataFrame(list(csv.reader(io.StringIO('\n'.join(lineas))))[1:], columns=['CodigoMed', 'NombreMed'])
    df['CodigoMed'] = df.CodigoMed.str.strip()
    df['NombreMed'] = df.NombreMed.map(lambda s: re.sub(r'[\x00-\x1f\x7f]', ' ', s).strip())
    return df.drop_duplicates('CodigoMed').set_index('CodigoMed').NombreMed


mi = leer_catalogo(ROOT / 'items/items_minsa.csv')
d = pd.read_excel(BASE, dtype=str).apply(lambda c: c.str.strip())
cols = list(d.columns)
N0 = len(d)
dup = pd.read_excel(ROOT / 'salida/dim_homologados_minsa_siga.xlsx', sheet_name='Duplicados_MINSA', dtype=str)
canon = dict(zip(dup.ID_SKU_EXTERNO, dup.CANONICO))
eliminados_dup = set(dup[dup.ACCION == 'ELIMINAR'].ID_SKU_EXTERNO)
val = pd.read_excel(ROOT / 'salida/validacion_manual_homologacion.xlsx', sheet_name=None, dtype=str)

# SKU de sutura elegibles: marca SD/CQ, o cualquier marca si se vende a MINSA
S = it[sut & (it.MARCA.isin(['SD', 'CQ']) | it.Coditem.isin(vend))].copy()
S['P'] = S.Item.map(V.parse_sku)
S['K'] = S.P.map(lambda p: (p['mat'], p['gauge'], p['needles']))
idx = S.groupby('K').Coditem.apply(list).to_dict()
PS = S.set_index('Coditem').P


def buscar_skus(desc_minsa):
    """SKU AUTO_SI para una descripción MINSA."""
    x = V.parse_minsa(desc_minsa)
    out = []
    for cod in idx.get((x['mat'], x['gauge'], x['needles']), []):
        r, m = V.evaluar(desc_minsa, I.Item[cod], x, PS[cod])
        if r == 'AUTO_SI':
            out.append(cod)
    return out


cat_sut = mi[mi.str.contains('SUTURA') | mi.str.contains(r'CATGUT|NAILON|POLIPROPILENO|SEDA NEGRA|POLIGLA|POLIDIOXANONA')]
cat_sut = cat_sut[~cat_sut.index.isin(eliminados_dup)]
CP = {c: V.parse_minsa(n) for c, n in cat_sut.items()}
cidx = {}
for c, p in CP.items():
    cidx.setdefault((p['mat'], p['gauge'], p['needles']), []).append(c)


def buscar_minsa(cod):
    y = PS.get(cod) if cod in PS.index else V.parse_sku(I.Item[cod])
    out = []
    for c in cidx.get((y['mat'], y['gauge'], y['needles']), []):
        r, m = V.evaluar(mi[c], I.Item[cod], CP[c], y)
        if r == 'AUTO_SI':
            out.append(c)
    return out


def fila(ext, desc, cod, origen):
    r = I.loc[cod]
    b = {c: None for c in cols}
    b.update(ID_SKU_EXTERNO=ext, DESC_PRODUCTO_EXTERNO=desc, FUENTE='MINSA', ID_SKU_UNICO_UNILENE=r.UNICO,
             ID_SKU_UNILENE=cod, AGRUPADOR=r.Agrupador, SUBAGRUPADOR='SUTURAS', VALIDACION='OK')
    b['CodigoSIGA'] = d.loc[d.ID_SKU_EXTERNO == ext, 'CodigoSIGA'].dropna().iloc[0] if (d.ID_SKU_EXTERNO == ext).any() and d.loc[d.ID_SKU_EXTERNO == ext, 'CodigoSIGA'].notna().any() else None
    b['_ORIGEN'] = origen
    return b


log = []          # todas las decisiones
borrar = []       # filas a eliminar
nuevas = []       # filas a agregar
pares = set(zip(d.ID_SKU_EXTERNO, d.ID_SKU_UNILENE))

# ------------------------------------------------------------------ 1) filas actuales con error de nº de agujas / medida
es_sut = d.ID_SKU_UNILENE.map(I.Agrupador).eq('01. SUTURAS')
for r in d[es_sut & d.ID_SKU_UNILENE.isin(I.index)].itertuples(index=False):
    res, mot = V.evaluar(r.DESC_PRODUCTO_EXTERNO, I.Item[r.ID_SKU_UNILENE])
    if res == 'AUTO_NO' and ('nº agujas' in mot or r.ID_SKU_EXTERNO == '30005'):
        borrar.append({**r._asdict(), 'MOTIVO': mot, 'VENDIDO_MINSA': 'SI' if r.ID_SKU_UNILENE in vend else 'NO'})
    log.append({'ETAPA': 'Tabla actual', 'ID_SKU_EXTERNO': r.ID_SKU_EXTERNO, 'DESC_PRODUCTO_EXTERNO': r.DESC_PRODUCTO_EXTERNO,
                'ID_SKU_UNILENE': r.ID_SKU_UNILENE, 'ITEM_UNILENE': I.Item[r.ID_SKU_UNILENE], 'RESULTADO': res, 'MOTIVO': mot,
                'ACCION': 'ELIMINAR' if any(b['ID_SKU_EXTERNO'] == r.ID_SKU_EXTERNO and b['ID_SKU_UNILENE'] == r.ID_SKU_UNILENE for b in borrar[-1:]) else
                ('ALERTA' if res == 'AUTO_NO' else '')})
B = pd.DataFrame(borrar)
pares_borrar = set(zip(B.ID_SKU_EXTERNO, B.ID_SKU_UNILENE)) if len(B) else set()
# reemplazo: homólogo correcto para esos códigos MINSA
for ext in sorted(set(B.ID_SKU_EXTERNO) if len(B) else []):
    desc = d.loc[d.ID_SKU_EXTERNO == ext, 'DESC_PRODUCTO_EXTERNO'].iloc[0]
    for cod in buscar_skus(desc):
        if (ext, cod) not in pares or (ext, cod) in pares_borrar:
            if (ext, cod) not in pares:
                nuevas.append(fila(ext, desc, cod, 'Reemplazo (nº agujas correcto)'))

# ------------------------------------------------------------------ 2) candidatos 1B (suturas)
b1 = val['1B_Candidatos_siga']
b1 = b1[b1.AGRUPADOR == '01. SUTURAS']
for r in b1.itertuples(index=False):
    res, mot = V.evaluar(r.DESC_PRODUCTO_EXTERNO, r.ITEM_UNILENE)
    log.append({'ETAPA': '1B Candidatos siga', 'ID_SKU_EXTERNO': r.ID_SKU_EXTERNO, 'DESC_PRODUCTO_EXTERNO': r.DESC_PRODUCTO_EXTERNO,
                'ID_SKU_UNILENE': r.ID_SKU_UNILENE, 'ITEM_UNILENE': r.ITEM_UNILENE, 'RESULTADO': res, 'MOTIVO': mot,
                'ACCION': {'AUTO_SI': 'AGREGAR', 'AUTO_NO': 'DESCARTAR'}.get(res, 'MANUAL')})
    if res == 'AUTO_SI' and r.ID_SKU_UNILENE in S.Coditem.values and (r.ID_SKU_EXTERNO, r.ID_SKU_UNILENE) not in pares:
        nuevas.append(fila(r.ID_SKU_EXTERNO, r.DESC_PRODUCTO_EXTERNO, r.ID_SKU_UNILENE, 'Candidato siga AUTO_SI'))
# códigos MINSA de 1B con algún AUTO_SI: completar con todos sus SKU AUTO_SI
for ext in {n['ID_SKU_EXTERNO'] for n in nuevas if n['_ORIGEN'] == 'Candidato siga AUTO_SI'}:
    desc = next(n['DESC_PRODUCTO_EXTERNO'] for n in nuevas if n['ID_SKU_EXTERNO'] == ext)
    for cod in buscar_skus(desc):
        if (ext, cod) not in pares:
            nuevas.append(fila(ext, desc, cod, 'Familia AUTO_SI'))

# ------------------------------------------------------------------ 3) vendidos sin MINSA (1C, suturas)
c1 = val['1C_Vendidos_sin_MINSA']
c1 = c1[c1.AGRUPADOR == '01. SUTURAS']
ext_dim = set(d.ID_SKU_EXTERNO)
for r in c1.itertuples(index=False):
    cods = buscar_minsa(r.ID_SKU_UNILENE)
    cods = sorted({canon.get(c, c) for c in cods}, key=lambda c: (c not in ext_dim, '*' in mi[c], c))
    res = 'AUTO_SI' if cods else 'MANUAL'
    log.append({'ETAPA': '1C Vendidos sin MINSA', 'ID_SKU_EXTERNO': ', '.join(cods) or None,
                'DESC_PRODUCTO_EXTERNO': mi[cods[0]] if cods else None, 'ID_SKU_UNILENE': r.ID_SKU_UNILENE,
                'ITEM_UNILENE': r.ITEM_UNILENE, 'RESULTADO': res,
                'MOTIVO': 'código MINSA con atributos iguales' if cods else 'sin código MINSA con atributos iguales',
                'ACCION': 'AGREGAR' if cods else 'MANUAL'})
    for c in cods:
        if (c, r.ID_SKU_UNILENE) not in pares:
            nuevas.append(fila(c, mi[c], r.ID_SKU_UNILENE, 'Vendido sin MINSA AUTO_SI'))

N = pd.DataFrame(nuevas).drop_duplicates(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE'])
N = N[~N.apply(lambda r: (r.ID_SKU_EXTERNO, r.ID_SKU_UNILENE) in pares, axis=1)]
N['VENDIDO_MINSA'] = np.where(N.ID_SKU_UNILENE.isin(vend), 'SI', 'NO')
N['ITEM_UNILENE'] = N.ID_SKU_UNILENE.map(I.Item)
LOG = pd.DataFrame(log)

# ------------------------------------------------------------------ nueva base
nd = d[~d.apply(lambda r: (r.ID_SKU_EXTERNO, r.ID_SKU_UNILENE) in pares_borrar, axis=1)]
nd = pd.concat([nd, N[cols]], ignore_index=True).sort_values(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE'])
assert not nd.duplicated(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE']).any()
wb = Workbook()
ws = wb.active
ws.title = 'sheet1'
ws.append(cols)
for r in nd[cols].itertuples(index=False):
    v = [None if x is None or (isinstance(x, float) and np.isnan(x)) else x for x in r]
    if str(v[0]).isdigit():
        v[0] = int(v[0])
    ws.append(v)
wb.save(SRC)

# ------------------------------------------------------------------ SQL parte 6
q = lambda x: 'NULL' if x is None or (isinstance(x, float) and np.isnan(x)) else "N'" + str(x).replace("'", "''") + "'"
L = ['-- DIM_HOMOLOGADOS – PARTE 6: validación automática de suturas (atributos.py)',
     f'-- Generado: {HOY} | Eliminar: {len(B)} filas | Insertar: {len(N)} filas',
     '-- Reglas: material, calibre, curvatura, punta, nº agujas, color y nº hebras iguales; hebra sin tolerancia; aguja ±0.5 mm.',
     '-- Requiere partes 1 a 5 ejecutadas. Compatible SQL Server 2012.', 'SET XACT_ABORT ON;', '',
     "IF OBJECT_ID('dbo.DIM_HOMOLOGADOS_AUTO_BK_20261007','U') IS NULL",
     "    SELECT * INTO dbo.DIM_HOMOLOGADOS_AUTO_BK_20261007 FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = N'MINSA';", 'GO', '',
     'BEGIN TRANSACTION;', "SELECT COUNT(*) AS filas_antes FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = N'MINSA';", '',
     f'-- 1) Eliminar {len(B)} filas con nº de agujas (o medida) distinto a la descripción MINSA']
for r in (B.itertuples() if len(B) else []):
    L.append(f"DELETE FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = N'MINSA' AND ID_SKU_EXTERNO = {q(r.ID_SKU_EXTERNO)} AND ID_SKU_UNILENE = {q(r.ID_SKU_UNILENE)};  -- {r.MOTIVO}")
L += ['', f'-- 2) Insertar {len(N)} filas AUTO_SI']
ic = cols
for org, g in N.groupby('_ORIGEN'):
    L.append(f'-- {org}: {len(g)} filas')
    for r in g.itertuples():
        L.append(f"IF NOT EXISTS (SELECT 1 FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = N'MINSA' AND ID_SKU_EXTERNO = {q(r.ID_SKU_EXTERNO)} AND ID_SKU_UNILENE = {q(r.ID_SKU_UNILENE)})\n"
                 f"    INSERT INTO dbo.DIM_HOMOLOGADOS ({', '.join(ic)}) VALUES ({', '.join(q(getattr(r, c)) for c in ic)});")
L += ['', "SELECT COUNT(*) AS filas_despues FROM dbo.DIM_HOMOLOGADOS WHERE FUENTE = N'MINSA';",
      f'-- Esperado: filas_antes - {len(B)} + {len(N)}  (Excel: {N0} -> {len(nd)})', '', '-- COMMIT TRANSACTION;', '-- ROLLBACK TRANSACTION;']
(ROOT / 'salida/validacion_auto_dim_homologados.sql').write_text('\n'.join(L) + '\n', encoding='utf-8-sig')

# ------------------------------------------------------------------ resultados
LOG.to_csv(ROOT / 'salida/decisiones_auto.csv', index=False, encoding='utf-8-sig')
cov0, cov1 = len(set(d.ID_SKU_UNILENE) & vend), len(set(nd.ID_SKU_UNILENE) & vend)
RES = pd.DataFrame([
    ('Filas antes', N0), ('Filas eliminadas (nº agujas / medida)', -len(B)),
    ('  de ellas con SKU vendido', int((B.VENDIDO_MINSA == 'SI').sum()) if len(B) else 0),
    ('Filas agregadas AUTO_SI', len(N))] + [(f'  {k}', v) for k, v in N._ORIGEN.value_counts().items()] + [
    ('Filas después', len(nd)), ('Códigos MINSA después', nd.ID_SKU_EXTERNO.nunique()),
    ('SKU vendidos cubiertos (antes → después)', f'{cov0} → {cov1} de {len(vend)}'), ('', ''),
    ('1B suturas: AUTO_SI / AUTO_NO / MANUAL', ' / '.join(str(int(((LOG.ETAPA == '1B Candidatos siga') & (LOG.RESULTADO == k)).sum())) for k in ('AUTO_SI', 'AUTO_NO', 'MANUAL'))),
    ('1C suturas: con código MINSA / sin código', ' / '.join(str(int(((LOG.ETAPA == '1C Vendidos sin MINSA') & (LOG.RESULTADO == k)).sum())) for k in ('AUTO_SI', 'MANUAL'))),
    ('Tabla actual: alertas AUTO_NO que se mantienen', int((LOG.ACCION == 'ALERTA').sum()))], columns=['Concepto', 'Valor'])
sheets = [('Resumen', RES), ('Eliminadas', B), ('Agregadas', N.rename(columns={'_ORIGEN': 'ORIGEN'})),
          ('Alertas_tabla_actual', LOG[LOG.ACCION == 'ALERTA']), ('Decisiones', LOG)]
HF = PatternFill('solid', fgColor='1F4E78')
COL = {'AUTO_SI': 'C6EFCE', 'AUTO_NO': 'FFC7CE', 'MANUAL': 'FFF2CC'}
with pd.ExcelWriter(ROOT / 'salida/validacion_auto_resultados.xlsx', engine='openpyxl') as w:
    for n, df in sheets:
        df.to_excel(w, sheet_name=n, index=False)
        ws = w.sheets[n]
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = HF
            c.alignment = Alignment(wrap_text=True, horizontal='center')
        for j, cn in enumerate(df.columns, 1):
            Lm = max([len(str(cn))] + [len(str(v)) for v in df[cn].head(400).fillna('')])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, Lm + 2), 60)
        if n != 'Resumen':
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        if 'RESULTADO' in df.columns:
            j = list(df.columns).index('RESULTADO') + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, j).fill = PatternFill('solid', fgColor=COL.get(ws.cell(r, j).value, 'FFFFFF'))
print(RES.to_string(index=False))
