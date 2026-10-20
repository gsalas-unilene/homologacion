"""Build salida/actualizar_dim_homologados.sql (sync of dbo.DIM_HOMOLOGADOS) from items/dim_homologados_minsa.xlsx.

Also validates the dim against the Items sheet of salida/estado_homologacion_items.xlsx.
Output is UTF-16 LE with BOM and CRLF so SSMS opens accents correctly.
"""
import collections
import datetime
import warnings

import openpyxl

warnings.simplefilter("ignore")


def t(v):
    return "" if v is None else str(v).strip()


dr = list(openpyxl.load_workbook("items/dim_homologados_minsa.xlsx", read_only=True).active.iter_rows(values_only=True))
H = list(dr[0])
dim = [dict(zip(H, r)) for r in dr[1:] if r[0] is not None]

# --- validation against the Items sheet
ws = openpyxl.load_workbook("salida/estado_homologacion_items.xlsx", read_only=True, data_only=True)["Items"]
rows = list(ws.iter_rows(min_row=6, values_only=True))
h = list(rows[0])
X = {t(r[0]): dict(zip(h, r)) for r in rows[1:] if r[0]}
dm = collections.defaultdict(lambda: {"min": set(), "siga": set(), "agr": set(), "sub": set()})
for r in dim:
    m = dm[t(r["ID_SKU_UNILENE"])]
    m["min"].add(t(r["ID_SKU_EXTERNO"]))
    if t(r["CodigoSIGA"]):
        m["siga"].add(t(r["CodigoSIGA"]))
    m["agr"].add(t(r["AGRUPADOR"]))
    m["sub"].add(t(r["SUBAGRUPADOR"]))
bad = collections.defaultdict(list)
for k, x in X.items():
    if t(x["HOMOLOGADO"]) != "SI":
        continue
    m = dm.get(k)
    if not m:
        bad["SKU homologado en Items y no en dim"].append(k)
        continue
    if {t(x["AGRUPADOR"])} != m["agr"]:
        bad["agrupador"].append(k)
    if {t(x["SUBAGRUPADOR"])} != m["sub"]:
        bad["subagrupador"].append(k)
    if set(t(x["CODIGOS_MINSA"]).split(" | ")) - {""} != m["min"]:
        bad["codigos MINSA"].append(k)
    if (set(t(x["CODIGOS_SIGA"]).split(" | ")) - {""}) != m["siga"] and len(m["min"]) == 1:
        bad["SIGA (SKU de un solo codigo)"].append(k)
for k in dm:
    if t(X.get(k, {}).get("HOMOLOGADO")) != "SI":
        bad["SKU en dim y no homologado en Items"].append(k)
print("filas dim", len(dim), "SKU", len(dm), "| discrepancias:", {a: v for a, v in bad.items()} or "ninguna")

# --- SQL
q = lambda s: "N'" + s.replace("'", "''") + "'"
COLS = ["ID_SKU_EXTERNO", "DESC_PRODUCTO_EXTERNO", "FUENTE", "ID_SKU_UNICO_UNILENE", "ID_SKU_UNILENE", "AGRUPADOR", "SUBAGRUPADOR",
        "AGRUPADOR_NEWCOMERCIAL", "SUB_AGRUPADOR_NEWCOMERCIAL", "LINEA_NEWCOMERCIAL", "DPRO_DESC_SUB_FAMILIA", "CodigoSIGA", "VALIDACION"]
vals = sorted(tuple(t(r[c]) for c in COLS) for r in dim)
val = lambda row: "  (" + ", ".join(q(v) if v else "NULL" for v in row) + ")"

out = [f"""/* Sync dbo.DIM_HOMOLOGADOS with the validated dimension items/dim_homologados_minsa.xlsx ({len(vals)} rows, {len(dm)} SKUs).
   SQL Server 2012. Generated {datetime.date.today()}. Only deletes the pair MINSA 21200 - PTATA00003 (rest: see the review queries in step 2c).
   Run salida/eliminar_marca_ve.sql first (brand VE).
   Steps: 1 load reference · 2 preview · 3 apply inside a transaction that ends in ROLLBACK (change to COMMIT after checking).
   Applied in step 3: insert missing MINSA<->SKU pairs, update AGRUPADOR / SUBAGRUPADOR / *_NEWCOMERCIAL where they differ,
   and CodigoSIGA / VALIDACION only if those columns exist in the table. */
SET NOCOUNT ON;
GO

IF OBJECT_ID('tempdb..#ref') IS NOT NULL DROP TABLE #ref;
CREATE TABLE #ref (
    ID_SKU_EXTERNO             VARCHAR(50)   NOT NULL,
    DESC_PRODUCTO_EXTERNO      NVARCHAR(300) NULL,
    FUENTE                     VARCHAR(20)   NOT NULL,
    ID_SKU_UNICO_UNILENE       VARCHAR(30)   NULL,
    ID_SKU_UNILENE             VARCHAR(20)   NOT NULL,
    AGRUPADOR                  NVARCHAR(100) NULL,
    SUBAGRUPADOR               NVARCHAR(150) NULL,
    AGRUPADOR_NEWCOMERCIAL     NVARCHAR(100) NULL,
    SUB_AGRUPADOR_NEWCOMERCIAL NVARCHAR(150) NULL,
    LINEA_NEWCOMERCIAL         NVARCHAR(100) NULL,
    DPRO_DESC_SUB_FAMILIA      NVARCHAR(150) NULL,
    CodigoSIGA                 VARCHAR(12)   NULL,
    VALIDACION                 VARCHAR(10)   NULL,
    PRIMARY KEY (FUENTE, ID_SKU_EXTERNO, ID_SKU_UNILENE)
);
GO
"""]
cols = ", ".join(COLS)
for i in range(0, len(vals), 500):
    out.append(f"INSERT INTO #ref ({cols}) VALUES\n" + ",\n".join(val(r) for r in vals[i:i + 500]) + ";\nGO\n")

out.append("""
/* 2) Preview ------------------------------------------------------------------------------------- */
-- 2a) pairs to insert
SELECT r.* FROM #ref r
WHERE NOT EXISTS (SELECT 1 FROM dbo.DIM_HOMOLOGADOS h
                  WHERE CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE);

-- 2b) rows to update (current vs new)
SELECT h.ID_SKU_EXTERNO, h.ID_SKU_UNILENE,
       h.AGRUPADOR AS AGRUPADOR_ACTUAL, r.AGRUPADOR AS AGRUPADOR_NUEVO,
       h.SUBAGRUPADOR AS SUBAGRUPADOR_ACTUAL, r.SUBAGRUPADOR AS SUBAGRUPADOR_NUEVO,
       h.LINEA_NEWCOMERCIAL AS LINEA_NC_ACTUAL, r.LINEA_NEWCOMERCIAL AS LINEA_NC_NUEVA
FROM dbo.DIM_HOMOLOGADOS h
JOIN #ref r ON CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE
WHERE ISNULL(h.AGRUPADOR, N'') <> ISNULL(r.AGRUPADOR, N'')
   OR ISNULL(h.SUBAGRUPADOR, N'') <> ISNULL(r.SUBAGRUPADOR, N'')
   OR ISNULL(h.AGRUPADOR_NEWCOMERCIAL, N'') <> ISNULL(r.AGRUPADOR_NEWCOMERCIAL, N'')
   OR ISNULL(h.SUB_AGRUPADOR_NEWCOMERCIAL, N'') <> ISNULL(r.SUB_AGRUPADOR_NEWCOMERCIAL, N'')
   OR ISNULL(h.LINEA_NEWCOMERCIAL, N'') <> ISNULL(r.LINEA_NEWCOMERCIAL, N'');

-- 2c) rows in the table that are not in the validated dimension (NOT deleted: review them)
SELECT h.* FROM dbo.DIM_HOMOLOGADOS h
WHERE NOT EXISTS (SELECT 1 FROM #ref r
                  WHERE r.ID_SKU_EXTERNO = CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) AND r.ID_SKU_UNILENE = h.ID_SKU_UNILENE);
GO

/* 3) Apply --------------------------------------------------------------------------------------- */
BEGIN TRANSACTION;

IF COL_LENGTH('dbo.DIM_HOMOLOGADOS', 'CodigoSIGA') IS NOT NULL AND COL_LENGTH('dbo.DIM_HOMOLOGADOS', 'VALIDACION') IS NOT NULL
    EXEC (N'INSERT INTO dbo.DIM_HOMOLOGADOS (ID_SKU_EXTERNO, DESC_PRODUCTO_EXTERNO, FUENTE, ID_SKU_UNICO_UNILENE, ID_SKU_UNILENE,
                AGRUPADOR, SUBAGRUPADOR, AGRUPADOR_NEWCOMERCIAL, SUB_AGRUPADOR_NEWCOMERCIAL, LINEA_NEWCOMERCIAL,
                DPRO_DESC_SUB_FAMILIA, CodigoSIGA, VALIDACION)
            SELECT r.ID_SKU_EXTERNO, r.DESC_PRODUCTO_EXTERNO, r.FUENTE, r.ID_SKU_UNICO_UNILENE, r.ID_SKU_UNILENE,
                r.AGRUPADOR, r.SUBAGRUPADOR, r.AGRUPADOR_NEWCOMERCIAL, r.SUB_AGRUPADOR_NEWCOMERCIAL, r.LINEA_NEWCOMERCIAL,
                r.DPRO_DESC_SUB_FAMILIA, r.CodigoSIGA, ISNULL(r.VALIDACION, ''OK'')
            FROM #ref r
            WHERE NOT EXISTS (SELECT 1 FROM dbo.DIM_HOMOLOGADOS h
                  WHERE CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE)');
ELSE
    INSERT INTO dbo.DIM_HOMOLOGADOS (ID_SKU_EXTERNO, DESC_PRODUCTO_EXTERNO, FUENTE, ID_SKU_UNICO_UNILENE, ID_SKU_UNILENE,
                AGRUPADOR, SUBAGRUPADOR, AGRUPADOR_NEWCOMERCIAL, SUB_AGRUPADOR_NEWCOMERCIAL, LINEA_NEWCOMERCIAL,
                DPRO_DESC_SUB_FAMILIA)
    SELECT r.ID_SKU_EXTERNO, r.DESC_PRODUCTO_EXTERNO, r.FUENTE, r.ID_SKU_UNICO_UNILENE, r.ID_SKU_UNILENE,
                r.AGRUPADOR, r.SUBAGRUPADOR, r.AGRUPADOR_NEWCOMERCIAL, r.SUB_AGRUPADOR_NEWCOMERCIAL, r.LINEA_NEWCOMERCIAL,
                r.DPRO_DESC_SUB_FAMILIA
    FROM #ref r
    WHERE NOT EXISTS (SELECT 1 FROM dbo.DIM_HOMOLOGADOS h
                  WHERE CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE);
PRINT CAST(@@ROWCOUNT AS VARCHAR(10)) + ' pares insertados';

/* Pairs removed on purpose (validated in estado_homologacion_items.xlsx) */
DELETE FROM dbo.DIM_HOMOLOGADOS
WHERE CAST(ID_SKU_EXTERNO AS VARCHAR(50)) = '21200' AND ID_SKU_UNILENE = 'PTATA00003';
PRINT CAST(@@ROWCOUNT AS VARCHAR(10)) + ' pares eliminados (esperado: 1)';

UPDATE h
SET h.AGRUPADOR = r.AGRUPADOR, h.SUBAGRUPADOR = r.SUBAGRUPADOR,
    h.AGRUPADOR_NEWCOMERCIAL = r.AGRUPADOR_NEWCOMERCIAL, h.SUB_AGRUPADOR_NEWCOMERCIAL = r.SUB_AGRUPADOR_NEWCOMERCIAL,
    h.LINEA_NEWCOMERCIAL = r.LINEA_NEWCOMERCIAL
FROM dbo.DIM_HOMOLOGADOS h
JOIN #ref r ON CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE
WHERE ISNULL(h.AGRUPADOR, N'') <> ISNULL(r.AGRUPADOR, N'')
   OR ISNULL(h.SUBAGRUPADOR, N'') <> ISNULL(r.SUBAGRUPADOR, N'')
   OR ISNULL(h.AGRUPADOR_NEWCOMERCIAL, N'') <> ISNULL(r.AGRUPADOR_NEWCOMERCIAL, N'')
   OR ISNULL(h.SUB_AGRUPADOR_NEWCOMERCIAL, N'') <> ISNULL(r.SUB_AGRUPADOR_NEWCOMERCIAL, N'')
   OR ISNULL(h.LINEA_NEWCOMERCIAL, N'') <> ISNULL(r.LINEA_NEWCOMERCIAL, N'');
PRINT CAST(@@ROWCOUNT AS VARCHAR(10)) + ' filas actualizadas (agrupador / subagrupador / newcomercial)';

IF COL_LENGTH('dbo.DIM_HOMOLOGADOS', 'CodigoSIGA') IS NOT NULL
    EXEC (N'UPDATE h SET h.CodigoSIGA = r.CodigoSIGA
            FROM dbo.DIM_HOMOLOGADOS h
            JOIN #ref r ON CAST(h.ID_SKU_EXTERNO AS VARCHAR(50)) = r.ID_SKU_EXTERNO AND h.ID_SKU_UNILENE = r.ID_SKU_UNILENE
            WHERE r.CodigoSIGA IS NOT NULL AND ISNULL(h.CodigoSIGA, '''') <> r.CodigoSIGA;
            PRINT CAST(@@ROWCOUNT AS VARCHAR(10)) + '' CodigoSIGA actualizados''');

ROLLBACK TRANSACTION;   -- change to COMMIT TRANSACTION once the preview is correct
GO

DROP TABLE #ref;
GO
""")
txt = "\n".join(out).replace("\r\n", "\n").replace("\n", "\r\n")
with open("salida/actualizar_dim_homologados.sql", "wb") as f:
    f.write(b"\xff\xfe" + txt.encode("utf-16-le"))
back = open("salida/actualizar_dim_homologados.sql", "rb").read().decode("utf-16").lstrip("﻿")
print("SQL: filas", back.count("\r\n  (N'"), "| roundtrip", back == txt, "| bytes", len(open("salida/actualizar_dim_homologados.sql", "rb").read()))
