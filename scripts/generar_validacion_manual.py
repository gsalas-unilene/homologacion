"""Libro de validación manual de la homologación MINSA (fase 1: homologación, fase 2: CodigoSIGA).

Ejecutar desde la raíz del repo:  python scripts/generar_validacion_manual.py
Base: items/dim_homologados_minsa.xlsx (= tabla DIM_HOMOLOGADOS, 2046 filas tras la parte 4).
Salida: salida/validacion_manual_homologacion.xlsx
"""
import csv
import datetime
import io
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rapidfuzz import fuzz, process
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import validador_auto as V

ROOT = Path('.')
HOY = datetime.date.today()
OUT = ROOT / 'salida/validacion_manual_homologacion.xlsx'
RX_DOBLE = re.compile(r'C/\s?2\s?A\b|DOBLE AGUJA|\b2C/A|C/2 AGUJAS')

# ------------------------------------------------------------------ fuentes
it = pd.read_csv(ROOT / 'items/items.csv', dtype=str).apply(lambda c: c.str.strip())
sut = (it.Agrupador == '01. SUTURAS') & (it.CodSut.str.len() == 21)
it['ES_SUTURA'] = sut
it['DOBLE'] = sut & (it.CodSut.str[17] == 'B')
it['UNICO'] = np.where(sut, it.CodSut.str[2:15], it.Coditem)          # convención actual de la tabla
it['MARCA'] = np.where(sut, it.CodSut.str[:2], '')
I = it.set_index('Coditem')


def leer_catalogo(path):
    lineas = []
    for raw in path.read_bytes().splitlines():
        try:
            lineas.append(raw.decode('utf-8'))
        except UnicodeDecodeError:
            lineas.append(raw.decode('cp1252', errors='replace'))
    df = pd.DataFrame(list(csv.reader(io.StringIO('\n'.join(lineas))))[1:], columns=['CodigoMed', 'NombreMed'])
    df['CodigoMed'] = df.CodigoMed.str.strip()
    df['NombreMed'] = df.NombreMed.map(lambda s: re.sub(r'[\x00-\x1f\x7f]', ' ', s).strip())
    return df.drop_duplicates('CodigoMed').set_index('CodigoMed').NombreMed


mi = leer_catalogo(ROOT / 'items/items_minsa.csv')
vend = set(pd.read_excel(ROOT / 'items/items_vendidos_minsa.xlsx', header=None, dtype=str)[0].str.strip())
sg = pd.read_excel(ROOT / 'salida/homologacion_siga_final.xlsx', dtype=str).apply(lambda c: c.str.strip()).set_index('CodigoMed')
d = pd.read_excel(ROOT / 'items/dim_homologados_minsa.xlsx', dtype=str).apply(lambda c: c.str.strip())
an = ROOT / 'salida/dim_homologados_minsa_siga.xlsx'
dup = pd.read_excel(an, sheet_name='Duplicados_MINSA', dtype=str)
canon = dict(zip(dup.ID_SKU_EXTERNO, dup.CANONICO))
pares = set(zip(d.ID_SKU_EXTERNO, d.ID_SKU_UNILENE))
_dec = ROOT / 'salida/decisiones_auto.csv'
DEC = pd.read_csv(_dec, dtype=str) if _dec.exists() else pd.DataFrame(columns=['ETAPA', 'ID_SKU_EXTERNO', 'ID_SKU_UNILENE', 'RESULTADO', 'MOTIVO'])
DEC1B = {(r.ID_SKU_EXTERNO, r.ID_SKU_UNILENE): (r.RESULTADO, r.MOTIVO) for r in DEC[DEC.ETAPA == '1B Candidatos siga'].itertuples()}
sku_dim = set(d.ID_SKU_UNILENE)
desc_minsa = {**sg.NombreMed.to_dict(), **mi.to_dict()}


def agr(cod, defecto=None):
    if cod in I.index:
        return I.Agrupador[cod]
    return defecto or 'SIN MAESTRO (no está en items.csv)'


# ------------------------------------------------------------------ F1A: filas actuales del dim a revisar
f1a = []
for r in d.itertuples(index=False):
    motivos = []
    if r.ID_SKU_UNILENE not in I.index:
        motivos.append('SKU no existe en items.csv')
    else:
        if I.Agrupador[r.ID_SKU_UNILENE] == '01. SUTURAS':
            res_, mot_ = V.evaluar(r.DESC_PRODUCTO_EXTERNO, I.Item[r.ID_SKU_UNILENE])
            if res_ == 'AUTO_NO':
                motivos.append('Atributos distintos: ' + mot_)
    if r.ID_SKU_EXTERNO == '30005' and r.ID_SKU_UNICO_UNILENE == 'CC030MR300075':
        motivos.append('Medida distinta: MINSA 40 mm, SKU 30 mm')
    if motivos:
        f1a.append({'AGRUPADOR': agr(r.ID_SKU_UNILENE, r.AGRUPADOR), 'MOTIVO': '; '.join(motivos),
                    'VENDIDO_MINSA': 'SI' if r.ID_SKU_UNILENE in vend else 'NO',
                    'ID_SKU_EXTERNO': r.ID_SKU_EXTERNO, 'DESC_PRODUCTO_EXTERNO': r.DESC_PRODUCTO_EXTERNO,
                    'ID_SKU_UNILENE': r.ID_SKU_UNILENE, 'ITEM_UNILENE': I.Item.get(r.ID_SKU_UNILENE),
                    'CodSut': I.CodSut.get(r.ID_SKU_UNILENE)})
F1A = pd.DataFrame(f1a)

# ------------------------------------------------------------------ F1B: candidatos de homologacion_siga_final
rm = pd.read_excel(an, sheet_name='Revision_manual', dtype=str)
rv = pd.read_excel(an, sheet_name='Revision_vendidos', dtype=str)
C = ['ID_SKU_EXTERNO', 'ID_SKU_UNILENE', 'TIER_COL_F', 'COL', 'PRIORIDAD', 'DIFERENCIAS', 'ITEM_C']
cand = pd.concat([rm.assign(COL=rm.COL_ORIGEN)[C], rv.assign(PRIORIDAD=None)[C]]).apply(
    lambda c: c.str.strip() if c.dtype == object else c)
cand['ID_SKU_EXTERNO'] = cand.ID_SKU_EXTERNO.map(lambda e: canon.get(e, e))
cand = cand[~cand.apply(lambda r: (r.ID_SKU_EXTERNO, r.ID_SKU_UNILENE) in pares, axis=1)]
cand = cand[cand.ID_SKU_UNILENE.isin(I.index)]
cand['PRIORIDAD'] = cand.PRIORIDAD.fillna(cand.COL.map({'C': '1', 'I': '2', 'J': '3'}))
cand = cand.sort_values('PRIORIDAD').drop_duplicates(['ID_SKU_EXTERNO', 'ID_SKU_UNILENE'])
cand['VALIDACION_AUTO'] = [DEC1B.get((a, b), ('', ''))[0] for a, b in zip(cand.ID_SKU_EXTERNO, cand.ID_SKU_UNILENE)]
cand['MOTIVO_AUTO'] = [DEC1B.get((a, b), ('', ''))[1] for a, b in zip(cand.ID_SKU_EXTERNO, cand.ID_SKU_UNILENE)]
DESCARTADOS = cand[cand.VALIDACION_AUTO == 'AUTO_NO'].copy()
cand = cand[~cand.VALIDACION_AUTO.isin(['AUTO_SI', 'AUTO_NO'])]
cand['DESC_PRODUCTO_EXTERNO'] = cand.ID_SKU_EXTERNO.map(desc_minsa)
cand['ITEM_UNILENE'] = cand.ID_SKU_UNILENE.map(I.Item)
cand['MARCA'] = cand.ID_SKU_UNILENE.map(I.MARCA)
cand['ID_SKU_UNICO_UNILENE'] = cand.ID_SKU_UNILENE.map(I.UNICO)
cand['AGRUPADOR'] = cand.ID_SKU_UNILENE.map(I.Agrupador)
cand['VENDIDO_MINSA'] = np.where(cand.ID_SKU_UNILENE.isin(vend), 'SI', 'NO')
cand['SKU_YA_HOMOLOGADO'] = np.where(cand.ID_SKU_UNILENE.isin(sku_dim), 'SI', 'NO')
cand['MINSA_YA_EN_DIM'] = np.where(cand.ID_SKU_EXTERNO.isin(set(d.ID_SKU_EXTERNO)), 'SI', 'NO')
cand['TIER_COL_F'] = cand.TIER_COL_F.fillna('')
# prioridad de trabajo: 1 vendido sin homologar, 2 vendido, 3 probable, 4 revisar
cand['ORDEN'] = np.select([(cand.VENDIDO_MINSA == 'SI') & (cand.SKU_YA_HOMOLOGADO == 'NO'), cand.VENDIDO_MINSA == 'SI',
                           cand.TIER_COL_F == 'probable'], [1, 2, 3], 4)
cand['NIVEL_TRABAJO'] = cand.ORDEN.map({1: 'A – vendido sin homologar', 2: 'B – vendido (ya homologado en otro MINSA)',
                                        3: 'C – probable', 4: 'D – revisar'})
F1B = cand.sort_values(['AGRUPADOR', 'ORDEN', 'ID_SKU_EXTERNO', 'PRIORIDAD'])

# ------------------------------------------------------------------ F1C: vendidos sin homologar y sin candidato -> sugerencias del catálogo MINSA
def norm(s):
    s = unicodedata.normalize('NFKD', str(s).upper()).encode('ascii', 'ignore').decode()
    rep = [(r'\bNYLON\b', 'NAILON'), (r'\bAGUJA\b', 'C/A'), (r'\bCIRCULO\b', 'CIRCULO'), (r'\bX UN\.?\b', ''),
           (r'\b(SD|CQ|CP|VE|VS|DM|HE)\s*$', ''), (r'\bPOLIGLACTIN\b', 'ACIDO POLIGLACTIN'), (r'\bMM\b', 'mm'),
           (r'[(),;]', ' ')]
    for a, b in rep:
        s = re.sub(a, b, s)
    return re.sub(r'\s+', ' ', s).strip()


MATERIALES = ['POLIGLACTIN', 'POLIGLICOLICO', 'POLIDIOXANONA', 'POLIPROPILENO', 'NAILON', 'SEDA', 'CATGUT',
              'CROMICO', 'SIMPLE', 'POLIESTER', 'ACERO', 'LINO', 'TEFLON', 'COLAGENO', 'POLIETILENO']
PUNTAS = ['REDONDA', 'CORTANTE', 'ESPATULADA', 'TROCAR']
nums = lambda s: set(re.findall(r'\d+(?:[./]\d+)?', str(s)))
cat_norm = {c: norm(n) for c, n in mi.items()}
cat_keys = list(cat_norm.keys())
cat_vals = [cat_norm[k] for k in cat_keys]
con_cand = set(F1B[F1B.VENDIDO_MINSA == 'SI'].ID_SKU_UNILENE)
falt = sorted(vend - sku_dim)
f1c, f1d = [], []
for sku in falt:
    if sku in con_cand:
        continue
    if sku not in I.index:
        f1d.append({'AGRUPADOR': 'SIN MAESTRO (no está en items.csv)', 'PREFIJO': sku[:5], 'ID_SKU_UNILENE': sku,
                    'ACCION': 'Completar maestro (Item, CodSut, Agrupador) y luego buscar código MINSA'})
        continue
    item = I.Item[sku]
    q = norm(item)
    pool = process.extract(q, cat_vals, scorer=fuzz.token_set_ratio, limit=40)
    ni = nums(item)

    def puntaje(x):
        nc = nums(cat_vals[x[2]])
        jac = len(ni & nc) / len(ni | nc) if (ni | nc) else 0
        cv = cat_vals[x[2]]
        pen = 0
        for grupo in (MATERIALES, PUNTAS):        # material y tipo de punta deben coincidir
            a, b = {w for w in grupo if w in q}, {w for w in grupo if w in cv}
            if a and a != b:
                pen += 30
        if ('C/2A' in q or 'DOBLE' in q) != bool(RX_DOBLE.search(cv)):
            pen += 30
        return x[1] * 0.6 + jac * 40 - pen      # similitud de texto + medidas - incompatibilidades
    res = [(t_, round(puntaje((t_, s_, i_))), i_) for t_, s_, i_ in sorted(pool, key=puntaje, reverse=True)[:3]]
    row = {'AGRUPADOR': I.Agrupador[sku], 'ID_SKU_UNILENE': sku, 'ITEM_UNILENE': item, 'CodSut': I.CodSut[sku],
           'MARCA': I.MARCA[sku]}
    for k, (txt, score, idx) in enumerate(res, 1):
        cod = cat_keys[idx]
        row[f'SUG{k}_CODIGO_MINSA'] = cod
        row[f'SUG{k}_DESCRIPCION'] = mi[cod]
        row[f'SUG{k}_SIMILITUD'] = max(0, min(100, round(score)))
        row[f'SUG{k}_MEDIDAS_IGUALES'] = 'SI' if nums(item) and nums(item) <= nums(mi[cod]) else 'NO'
        row[f'SUG{k}_YA_EN_DIM'] = 'SI' if canon.get(cod, cod) in set(d.ID_SKU_EXTERNO) else 'NO'
    f1c.append(row)
F1C = pd.DataFrame(f1c)
F1C['CONFIANZA'] = np.select([(F1C.SUG1_SIMILITUD >= 85) & (F1C.SUG1_MEDIDAS_IGUALES == 'SI'),
                              (F1C.SUG1_SIMILITUD >= 70)], ['ALTA', 'MEDIA'], 'BAJA')
F1C = F1C.sort_values(['AGRUPADOR', 'CONFIANZA', 'SUG1_SIMILITUD'], ascending=[True, True, False])
F1D = pd.DataFrame(f1d)

# ------------------------------------------------------------------ F1E: códigos MINSA que quedaron sin homólogo tras la validación automática
_b5 = ROOT / 'items/backup/dim_homologados_minsa_post_script5_20261007.xlsx'
f1e = []
if _b5.exists():
    d5 = pd.read_excel(_b5, dtype=str).apply(lambda c: c.str.strip())
    S_ = it[(it.Agrupador == '01. SUTURAS') & (it.CodSut.str.len() == 21) & (it.MARCA.isin(['SD', 'CQ']) | it.Coditem.isin(vend))]
    for ext in sorted(set(d5.ID_SKU_EXTERNO) - set(d.ID_SKU_EXTERNO)):
        desc = d5.loc[d5.ID_SKU_EXTERNO == ext, 'DESC_PRODUCTO_EXTERNO'].iloc[0]
        x = V.parse_minsa(desc)
        pool = S_[S_.Item.map(lambda s: (lambda y: y['mat'] == x['mat'] and y['gauge'] == x['gauge'])(V.parse_sku(s)))]
        ranked = sorted(((V.evaluar(desc, r.Item), r.Coditem, r.Item) for r in pool.itertuples()),
                        key=lambda z: (z[0][0] != 'AUTO_SI', z[0][1].count(';'), z[0][1]))[:3]
        row = {'AGRUPADOR': '01. SUTURAS', 'ID_SKU_EXTERNO': ext, 'DESC_PRODUCTO_EXTERNO': desc,
               'SKU_ANTERIORES (eliminados)': ', '.join(d5.loc[d5.ID_SKU_EXTERNO == ext, 'ID_SKU_UNILENE'])}
        for k, ((res_, mot_), cod, item) in enumerate(ranked, 1):
            row[f'SUG{k}_SKU'] = cod
            row[f'SUG{k}_ITEM'] = item
            row[f'SUG{k}_DIFERENCIA'] = mot_
        f1e.append(row)
F1E = pd.DataFrame(f1e)

# ------------------------------------------------------------------ F2: CodigoSIGA faltante (tarea final)
p = d.drop_duplicates('ID_SKU_EXTERNO').copy()
p = p[p.CodigoSIGA.isna()]
ag = d.assign(A=d.ID_SKU_UNILENE.map(I.Agrupador).fillna(d.AGRUPADOR)).groupby('ID_SKU_EXTERNO').A.first()
ven = d.assign(v=d.ID_SKU_UNILENE.isin(vend)).groupby('ID_SKU_EXTERNO').v.any()
F2 = pd.DataFrame({
    'AGRUPADOR': p.ID_SKU_EXTERNO.map(ag), 'ID_SKU_EXTERNO': p.ID_SKU_EXTERNO, 'DESC_PRODUCTO_EXTERNO': p.DESC_PRODUCTO_EXTERNO,
    'TIENE_VENTA': p.ID_SKU_EXTERNO.map(ven).map({True: 'SI', False: 'NO'}),
    'SIGA_ESTADO': p.ID_SKU_EXTERNO.map(sg.SIGA_estado).fillna('sin_busqueda'),
    'SIGA_CODIGO_SUGERIDO': p.ID_SKU_EXTERNO.map(sg.SIGA_codigo_1), 'SIGA_DESCRIPCION_SUGERIDA': p.ID_SKU_EXTERNO.map(sg.SIGA_descripcion_1),
    'SIGA_DIFERENCIAS': p.ID_SKU_EXTERNO.map(sg.SIGA_diferencias), 'SIGA_ALTERNATIVAS': p.ID_SKU_EXTERNO.map(sg.SIGA_alternativas)})
F2['TAREA'] = np.where(F2.SIGA_ESTADO == 'sin_busqueda', 'Buscar en SIGA', 'Confirmar sugerencia o elegir alternativa')
F2 = F2.sort_values(['AGRUPADOR', 'TIENE_VENTA', 'SIGA_ESTADO', 'ID_SKU_EXTERNO'], ascending=[True, False, True, True])

# ------------------------------------------------------------------ resumen por agrupador
def cnt(df, name, key=None):
    if df.empty:
        return pd.Series(dtype=int, name=name)
    return (df.groupby('AGRUPADOR')[key].nunique() if key else df.groupby('AGRUPADOR').size()).rename(name)


RES = pd.concat([
    cnt(F1A, '1A filas dim a revisar'),
    cnt(F1B[F1B.ORDEN == 1], '1B nivel A (vendido sin homologar)'),
    cnt(F1B, '1B candidatos (filas)'), cnt(F1B, '1B códigos MINSA', 'ID_SKU_EXTERNO'),
    cnt(F1C, '1C vendidos sin candidato'), cnt(F1D, '1D sin maestro'), cnt(F1E, '1E MINSA sin homólogo'),
    cnt(F2, '2 CodigoSIGA faltante (códigos)')], axis=1).fillna(0).astype(int)
RES.loc['TOTAL'] = RES.sum()
RES = RES.reset_index().rename(columns={'index': 'AGRUPADOR'})

PLAN = pd.DataFrame([
    ('FASE 1 – HOMOLOGACIÓN MINSA ↔ SKU (primero)', '', ''),
    ('Auto', 'validacion_auto_resultados.xlsx', 'Suturas ya validadas automáticamente (atributos iguales = agregadas; distintos = descartadas). Aquí solo queda lo MANUAL.'),
    ('1A', 'Revisar_dim_actual', 'Filas que YA están en DIM_HOMOLOGADOS con una señal de error (doble aguja, medida, SKU sin maestro). '
     'Marque DECISION: MANTENER o ELIMINAR.'),
    ('1B', 'Candidatos_siga', 'Propuestas del análisis siga_final (F = probable/revisar y alternativas I/J). Trabaje por AGRUPADOR y en orden '
     'NIVEL_TRABAJO A→D. Compare DESC_PRODUCTO_EXTERNO vs ITEM_UNILENE: material, calibre, tipo y medida de aguja, '
     'nº de agujas, largo de hebra. Marque DECISION: SI / NO.'),
    ('1C', 'Vendidos_sin_MINSA', 'SKU que se venden a MINSA y no tienen homologación ni candidato. Se proponen 3 códigos del catálogo MINSA '
     'por similitud. Escriba el código correcto en CODIGO_MINSA_ELEGIDO (o NINGUNO). Empiece por CONFIANZA ALTA.'),
    ('1E', 'MINSA_sin_homologo', 'Códigos MINSA cuyas filas se eliminaron por nº de agujas y no tienen SKU con atributos iguales. Se proponen 3 SKU con la diferencia. Escriba SKU_ELEGIDO o NINGUNO.'),
    ('1D', 'Sin_maestro', 'SKU vendidos que no existen en items.csv (65 DMNST). Primero completar el maestro; luego pasan a 1C.'),
    ('', '', ''),
    ('FASE 2 – CÓDIGO SIGA (al final)', '', ''),
    ('2', 'CodigoSIGA_faltante', 'Códigos MINSA del dim sin CodigoSIGA. Confirme la sugerencia o busque en SIGA; escriba CodigoSIGA_ELEGIDO. '
     'Hágalo después de cerrar la fase 1, porque los nuevos códigos MINSA aprobados también necesitarán SIGA.'),
    ('', '', ''),
    ('CÓMO DEVOLVERLO', '', ''),
    ('', '', 'Guarde este archivo con sus decisiones. Con él se generan automáticamente: INSERT (SI / código elegido), '
     'DELETE (ELIMINAR) y UPDATE de CodigoSIGA, en el mismo formato de los scripts anteriores.'),
    ('', '', 'Puede entregarlo por partes (por ejemplo, un agrupador a la vez).'),
    ('', '', ''),
    ('REGLAS DE DECISIÓN', '', ''),
    ('Exacto', '', 'Mismo material, calibre, tipo y medida de aguja, nº de agujas y largo de hebra → SI.'),
    ('Tolerancia', '', 'Diferencia solo en largo de hebra (ej. 70 vs 75 cm) → decisión comercial; anotar en OBSERVACION.'),
    ('No', '', 'Cambia calibre, tipo de aguja (redonda/cortante), curvatura (1/2 vs 3/8) o nº de agujas → NO.'),
    ('Marcas', '', 'Se mantienen solo SD y CQ, salvo SKU vendidos de otra marca con el mismo ID único.'),
], columns=['Paso', 'Hoja', 'Qué hacer'])

# ------------------------------------------------------------------ escritura
def col_order(df, first):
    return df[[c for c in first if c in df.columns] + [c for c in df.columns if c not in first]]


F1A = col_order(F1A.assign(DECISION='', OBSERVACION=''), ['DECISION', 'AGRUPADOR', 'VENDIDO_MINSA', 'MOTIVO', 'ID_SKU_EXTERNO',
                                                          'DESC_PRODUCTO_EXTERNO', 'ID_SKU_UNILENE', 'ITEM_UNILENE', 'CodSut', 'OBSERVACION'])
F1A = F1A.sort_values(['AGRUPADOR', 'VENDIDO_MINSA', 'ID_SKU_EXTERNO'], ascending=[True, False, True])
F1B = col_order(F1B.assign(DECISION='', OBSERVACION=''), ['DECISION', 'AGRUPADOR', 'NIVEL_TRABAJO', 'ID_SKU_EXTERNO', 'DESC_PRODUCTO_EXTERNO',
                                                          'ID_SKU_UNILENE', 'ITEM_UNILENE', 'MARCA', 'TIER_COL_F', 'COL', 'PRIORIDAD',
                                                          'DIFERENCIAS', 'MOTIVO_AUTO', 'VENDIDO_MINSA', 'SKU_YA_HOMOLOGADO', 'MINSA_YA_EN_DIM',
                                                          'ID_SKU_UNICO_UNILENE', 'ITEM_C', 'OBSERVACION']).drop(columns=['ORDEN', 'VALIDACION_AUTO'])
F1C = col_order(F1C.assign(CODIGO_MINSA_ELEGIDO='', OBSERVACION=''), ['CODIGO_MINSA_ELEGIDO', 'AGRUPADOR', 'CONFIANZA', 'ID_SKU_UNILENE', 'ITEM_UNILENE'])
F1D = col_order(F1D.assign(OBSERVACION=''), ['AGRUPADOR', 'PREFIJO', 'ID_SKU_UNILENE', 'ACCION'])
F1E = col_order(F1E.assign(SKU_ELEGIDO='', OBSERVACION=''), ['SKU_ELEGIDO', 'AGRUPADOR', 'ID_SKU_EXTERNO', 'DESC_PRODUCTO_EXTERNO'])
F2 = col_order(F2.assign(CodigoSIGA_ELEGIDO='', OBSERVACION=''), ['CodigoSIGA_ELEGIDO', 'AGRUPADOR', 'TAREA', 'TIENE_VENTA', 'ID_SKU_EXTERNO'])

sheets = [('0_PLAN', PLAN), ('0_RESUMEN_AGRUPADOR', RES), ('1A_Revisar_dim_actual', F1A), ('1B_Candidatos_siga', F1B),
          ('1C_Vendidos_sin_MINSA', F1C), ('1D_Sin_maestro', F1D), ('1E_MINSA_sin_homologo', F1E), ('2_CodigoSIGA_faltante', F2)]
listas = {'1A_Revisar_dim_actual': ('DECISION', '"MANTENER,ELIMINAR"'), '1B_Candidatos_siga': ('DECISION', '"SI,NO"')}
entrada = {'DECISION', 'CODIGO_MINSA_ELEGIDO', 'SKU_ELEGIDO', 'CodigoSIGA_ELEGIDO', 'OBSERVACION'}
HF = PatternFill('solid', fgColor='1F4E78')
YEL = PatternFill('solid', fgColor='FFF2CC')
SEC = PatternFill('solid', fgColor='D9E1F2')
BAND = PatternFill('solid', fgColor='EEF3F8')
thin = Border(top=Side(style='thin', color='9BB0C9'))
conf = {'ALTA': 'C6EFCE', 'MEDIA': 'FFEB9C', 'BAJA': 'F8CBAD'}

wb = Workbook()
wb.remove(wb.active)
for name, df in sheets:
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (v is None or (isinstance(v, float) and np.isnan(v)) or v == '') else
                   (int(v) if isinstance(v, (np.integer,)) else v) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HF
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 34
    for j, cn in enumerate(df.columns, 1):
        L = max([len(str(cn))] + [len(str(v)) for v in df[cn].head(600).fillna('')])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, L + 2), 110 if name == '0_PLAN' else 55)
    if name == '0_PLAN':
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical='top')
            if str(ws.cell(r, 1).value or '').isupper() and len(str(ws.cell(r, 1).value)) > 6:
                for j in (1, 2, 3):
                    ws.cell(r, j).fill = SEC
                    ws.cell(r, j).font = Font(bold=True, color='1F4E78')
        continue
    if name == '0_RESUMEN_AGRUPADOR':
        for j in range(1, ws.max_column + 1):
            ws.cell(ws.max_row, j).font = Font(bold=True)
            ws.cell(ws.max_row, j).fill = SEC
        continue
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    cols = list(df.columns)
    for cn in entrada & set(cols):
        j = cols.index(cn) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(r, j).fill = YEL
    if name in listas:
        cn, f = listas[name]
        dv = DataValidation(type='list', formula1=f, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'A2:A{ws.max_row}')
        ws.conditional_formatting.add(f'A2:A{ws.max_row}', FormulaRule(formula=['OR($A2="SI",$A2="MANTENER")'], fill=PatternFill('solid', fgColor='C6EFCE')))
        ws.conditional_formatting.add(f'A2:A{ws.max_row}', FormulaRule(formula=['OR($A2="NO",$A2="ELIMINAR")'], fill=PatternFill('solid', fgColor='FFC7CE')))
    # línea separadora al cambiar de agrupador y bandas por código MINSA
    if 'AGRUPADOR' in cols:
        ja = cols.index('AGRUPADOR') + 1
        jk = cols.index('ID_SKU_EXTERNO') + 1 if 'ID_SKU_EXTERNO' in cols else None
        prev, last, on = None, None, False
        for r in range(2, ws.max_row + 1):
            a = ws.cell(r, ja).value
            if a != prev and prev is not None:
                for j in range(1, ws.max_column + 1):
                    ws.cell(r, j).border = thin
            prev = a
            if jk:
                k = ws.cell(r, jk).value
                if k != last:
                    on, last = not on, k
                if on:
                    for j in range(ja, ws.max_column + 1):
                        if cols[j - 1] not in entrada:
                            ws.cell(r, j).fill = BAND
    if 'CONFIANZA' in cols:
        j = cols.index('CONFIANZA') + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(r, j).fill = PatternFill('solid', fgColor=conf.get(ws.cell(r, j).value, 'FFFFFF'))
wb.save(OUT)

print(RES.to_string(index=False))
print('F1A', F1A.MOTIVO.value_counts().to_dict(), 'vendidos', (F1A.VENDIDO_MINSA == 'SI').sum())
print('F1B niveles', F1B.NIVEL_TRABAJO.value_counts().to_dict())
print('F1C confianza', F1C.CONFIANZA.value_counts().to_dict())
print('F2', F2.TAREA.value_counts().to_dict(), 'con venta', (F2.TIENE_VENTA == 'SI').sum())
print(F1C[['ITEM_UNILENE', 'SUG1_CODIGO_MINSA', 'SUG1_DESCRIPCION', 'SUG1_SIMILITUD', 'CONFIANZA']].head(12).to_string())
